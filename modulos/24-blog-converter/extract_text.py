#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Blog Converter (Modulo 24) - Extrator de texto.

Recebe o caminho de um arquivo .docx ou .pdf e imprime no stdout um JSON:
    {"ok": true, "text": "...", "chars": 1234, "source": "docx|pdf"}
ou, em caso de erro:
    {"ok": false, "error": "mensagem"}

Uso:
    python extract_text.py <caminho-do-arquivo>

Dependencias: python-docx, pypdf (ja disponiveis no ambiente local do Office).
"""
import sys
import os
import json


def extract_docx(path):
    from docx import Document
    doc = Document(path)
    parts = []
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t:
            parts.append(t)
    # tabelas viram linhas separadas por " | "
    for table in doc.tables:
        for row in table.rows:
            cells = [(c.text or "").strip() for c in row.cells]
            line = " | ".join([c for c in cells if c])
            if line:
                parts.append(line)
    return "\n\n".join(parts)


def extract_pdf(path):
    from pypdf import PdfReader
    reader = PdfReader(path)
    parts = []
    for page in reader.pages:
        t = page.extract_text() or ""
        t = t.strip()
        if t:
            parts.append(t)
    return "\n\n".join(parts)


def main():
    if len(sys.argv) < 2:
        print(json.dumps({"ok": False, "error": "Caminho do arquivo nao informado."}))
        return

    path = sys.argv[1]
    if not os.path.isfile(path):
        print(json.dumps({"ok": False, "error": "Arquivo nao encontrado."}))
        return

    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".docx":
            text = extract_docx(path)
            source = "docx"
        elif ext == ".pdf":
            text = extract_pdf(path)
            source = "pdf"
        else:
            print(json.dumps({"ok": False, "error": "Formato nao suportado (use .docx ou .pdf)."}))
            return

        text = (text or "").strip()
        if not text:
            print(json.dumps({"ok": False, "error": "Nao foi possivel extrair texto do arquivo (documento vazio ou apenas imagens)."}))
            return

        print(json.dumps({"ok": True, "text": text, "chars": len(text), "source": source}, ensure_ascii=False))
    except Exception as e:  # noqa: BLE001
        print(json.dumps({"ok": False, "error": "Falha na extracao: %s" % str(e)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
